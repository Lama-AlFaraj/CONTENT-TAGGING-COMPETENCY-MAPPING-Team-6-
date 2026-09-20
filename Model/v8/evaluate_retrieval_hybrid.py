import os
import re
import numpy as np
import pandas as pd
import torch
from transformers import AutoTokenizer, AutoModel
from pptx import Presentation
from openpyxl import load_workbook
import nbformat

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "../.."))

TAXONOMY_PATH = os.path.join(ROOT, "Model/saudi_skills_taxonomy_v1_final.csv")
GOLD_PATH = "/tmp/competency_gold_labels_review.csv"

MODEL_NAME = "intfloat/multilingual-e5-base"

CHUNK_SIZE = 6000
OVERLAP = 500
TOP_K = 10

WEIGHTS = [1.0, 0.9, 0.8, 0.7, 0.6, 0.5]


def normalize_text(x):
    x = str(x).lower()
    x = re.sub(r"[^a-z0-9\u0600-\u06ff\s]", " ", x)
    return re.sub(r"\s+", " ", x).strip()


def chunk_text(text):
    text = str(text)
    chunks = []
    start = 0

    while start < len(text):
        end = min(len(text), start + CHUNK_SIZE)
        chunk = text[start:end]

        if chunk.strip():
            chunks.append(chunk)

        if end >= len(text):
            break

        start = end - OVERLAP

    return chunks


def extract_text(path):
    ext = os.path.splitext(path)[1].lower()

    if ext == ".pptx":
        prs = Presentation(path)
        parts = []

        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    parts.append(shape.text)

        return "\n".join(parts)

    if ext == ".xlsx":
        wb = load_workbook(path, read_only=True, data_only=True)
        parts = []

        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                vals = [str(v) for v in row if v is not None]
                if vals:
                    parts.append(" ".join(vals))

        return "\n".join(parts)

    if ext == ".ipynb":
        nb = nbformat.read(path, as_version=4)
        parts = []

        for cell in nb.cells:
            if cell.cell_type in ("markdown", "code"):
                parts.append(cell.source)

        return "\n".join(parts)

    if ext in (".md", ".txt", ".csv"):
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()

    return ""


def mean_pool(last_hidden_state, attention_mask):
    mask = attention_mask.unsqueeze(-1).expand(last_hidden_state.size()).float()
    summed = torch.sum(last_hidden_state * mask, dim=1)
    counts = torch.clamp(mask.sum(dim=1), min=1e-9)
    return summed / counts


print("=" * 80)
print("HYBRID RETRIEVAL EXPERIMENT")
print("=" * 80)

taxonomy = pd.read_csv(TAXONOMY_PATH)

required = {"skill_name_en", "description_en"}
missing = required - set(taxonomy.columns)

if missing:
    raise ValueError(f"Missing taxonomy columns: {missing}")

taxonomy["search_text"] = (
    taxonomy["skill_name_en"].fillna("").astype(str)
    + ". "
    + taxonomy["description_en"].fillna("").astype(str)
)

taxonomy["lexical_text"] = taxonomy["search_text"].map(normalize_text)

gold = pd.read_csv(GOLD_PATH)

print(f"Taxonomy skills: {len(taxonomy)}")
print(f"Gold rows: {len(gold)}")

tokenizer = AutoTokenizer.from_pretrained(MODEL_NAME)
model = AutoModel.from_pretrained(MODEL_NAME)

device = "cuda" if torch.cuda.is_available() else "cpu"
model = model.to(device)
model.eval()

print(f"Device: {device}")


@torch.no_grad()
def encode(texts):
    inputs = tokenizer(
        texts,
        padding=True,
        truncation=True,
        max_length=512,
        return_tensors="pt",
    )

    inputs = {k: v.to(device) for k, v in inputs.items()}

    outputs = model(**inputs)

    embeddings = mean_pool(
        outputs.last_hidden_state,
        inputs["attention_mask"],
    )

    embeddings = torch.nn.functional.normalize(embeddings, p=2, dim=1)

    return embeddings.cpu().numpy()


print("Encoding taxonomy...")

taxonomy_embeddings = encode(
    ["passage: " + x for x in taxonomy["search_text"].tolist()]
)

print("Taxonomy embeddings ready.")


def lexical_score(query, taxonomy_text):
    q_tokens = set(normalize_text(query).split())
    t_tokens = set(taxonomy_text.split())

    if not q_tokens or not t_tokens:
        return 0.0

    return len(q_tokens & t_tokens) / len(q_tokens)


def retrieve(query, weight, top_k=TOP_K):
    query_embedding = encode(["query: " + query])[0]

    e5_scores = taxonomy_embeddings @ query_embedding

    lexical_scores = np.array([
        lexical_score(query, text)
        for text in taxonomy["lexical_text"]
    ])

    # Normalize E5 to 0-1
    e5_min = e5_scores.min()
    e5_max = e5_scores.max()

    if e5_max > e5_min:
        e5_norm = (e5_scores - e5_min) / (e5_max - e5_min)
    else:
        e5_norm = np.zeros_like(e5_scores)

    # Lexical score is already 0-1
    hybrid_scores = (
        weight * e5_norm
        + (1 - weight) * lexical_scores
    )

    order = np.argsort(-hybrid_scores)[:top_k]

    return order, hybrid_scores


files = sorted(gold["file_name"].unique())

results = []

for file_name in files:
    print("\n" + "-" * 80)
    print(file_name)

    file_rows = gold[gold["file_name"] == file_name]

    gold_skills = set(
        file_rows["taxonomy_skill"]
        .dropna()
        .astype(str)
        .str.strip()
    )

    candidate_path = None

    for base in [
        os.path.join(ROOT, "Model/evaluation_data"),
        os.path.join(ROOT, "data"),
    ]:
        p = os.path.join(base, file_name)
        if os.path.exists(p):
            candidate_path = p
            break

    if candidate_path is None:
        print("WARNING: file not found")
        continue

    text = extract_text(candidate_path)
    chunks = chunk_text(text)

    print(f"Chunks: {len(chunks)}")
    print(f"Gold skills: {sorted(gold_skills)}")

    # Use several representative chunks.
    queries = chunks[:8]

    for weight in WEIGHTS:
        all_scores = []

        for query in queries:
            _, scores = retrieve(query, weight, top_k=len(taxonomy))
            all_scores.append(scores)

        all_scores = np.vstack(all_scores)

        # Max score across chunks for each skill.
        final_scores = all_scores.max(axis=0)

        order = np.argsort(-final_scores)[:TOP_K]

        predicted = set(
            taxonomy.iloc[order]["skill_name_en"].astype(str)
        )

        hits = len(gold_skills & predicted)

        recall = (
            hits / len(gold_skills)
            if gold_skills
            else 0.0
        )

        hit = 1 if hits > 0 else 0

        results.append({
            "file_name": file_name,
            "weight_e5": weight,
            "weight_lexical": round(1 - weight, 2),
            "gold_count": len(gold_skills),
            "hits": hits,
            "recall_at_10": recall,
            "hit_at_10": hit,
            "top10": " | ".join(
                taxonomy.iloc[order]["skill_name_en"].astype(str)
            ),
        })

        print(
            f"E5={weight:.1f} "
            f"Lexical={1-weight:.1f} "
            f"Recall@10={recall:.3f} "
            f"Hit={hit}"
        )


df = pd.DataFrame(results)

summary = (
    df.groupby(
        ["weight_e5", "weight_lexical"],
        as_index=False
    )
    .agg(
        mean_recall_at_10=("recall_at_10", "mean"),
        hit_rate_at_10=("hit_at_10", "mean"),
    )
)

summary["mean_recall_at_10_pct"] = (
    summary["mean_recall_at_10"] * 100
)

summary["hit_rate_at_10_pct"] = (
    summary["hit_rate_at_10"] * 100
)

summary = summary.sort_values(
    "mean_recall_at_10",
    ascending=False,
)

out_path = os.path.join(
    ROOT,
    "Model/v8/hybrid_retrieval_results.csv"
)

summary_path = os.path.join(
    ROOT,
    "Model/v8/hybrid_retrieval_summary.csv"
)

df.to_csv(out_path, index=False)
summary.to_csv(summary_path, index=False)

print("\n" + "=" * 80)
print("FINAL HYBRID RESULTS")
print("=" * 80)

print(
    summary[
        [
            "weight_e5",
            "weight_lexical",
            "mean_recall_at_10_pct",
            "hit_rate_at_10_pct",
        ]
    ].to_string(index=False)
)

print(f"\nSaved detailed results: {out_path}")
print(f"Saved summary: {summary_path}")
