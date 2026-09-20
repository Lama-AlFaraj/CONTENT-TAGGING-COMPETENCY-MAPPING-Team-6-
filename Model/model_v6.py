# -*- coding: utf-8 -*-

"""
AIDC BeamData Capstone Project - V6
Model: Qwen/Qwen2.5-7B-Instruct
Quantization: 4-bit NF4
Purpose: Content Tagging & Competency Mapping

V6 preserved:
- Qwen2.5-7B-Instruct
- 4-bit NF4 quantization
- content extraction
- 6000-character chunks
- 500-character overlap
- chunk-level prediction
- final aggregation
- 3-8 tags
- 1-3 taxonomy competencies
- difficulty classification
- confidence
- notes

Required additions:
- frozen Saudi Skills Taxonomy
- learning objectives
- structured final output
"""

import os
import json
import gc
from pathlib import Path

import torch
import pandas as pd

from transformers import (
    AutoTokenizer,
    AutoModelForCausalLM,
    BitsAndBytesConfig
)

from pptx import Presentation
from openpyxl import load_workbook


# ============================================================
# 1. MODEL SETUP
# ============================================================

MODEL_NAME = "Qwen/Qwen2.5-7B-Instruct"

print("=" * 80)
print("V6 MODEL SETUP")
print("=" * 80)

print("PyTorch version:", torch.__version__)
print("CUDA available:", torch.cuda.is_available())

if not torch.cuda.is_available():
    raise RuntimeError("CUDA is required for the V6 Qwen 7B model.")

print("GPU:", torch.cuda.get_device_name(0))


quant_config = BitsAndBytesConfig(
    load_in_4bit=True,
    bnb_4bit_compute_dtype=torch.float16,
    bnb_4bit_quant_type="nf4",
    bnb_4bit_use_double_quant=True
)

print("Loading tokenizer...")

tokenizer = AutoTokenizer.from_pretrained(
    MODEL_NAME
)

print("Loading Qwen 2.5 model...")

model = AutoModelForCausalLM.from_pretrained(
    MODEL_NAME,
    quantization_config=quant_config,
    device_map="auto"
)

print("Qwen 2.5 loaded successfully!")


# ============================================================
# 2. FROZEN SAUDI SKILLS TAXONOMY
# ============================================================

TAXONOMY_FILE = "saudi_skills_taxonomy_v1_final.csv"

if not os.path.exists(TAXONOMY_FILE):
    raise FileNotFoundError(
        f"Required taxonomy file not found: {TAXONOMY_FILE}"
    )

taxonomy = pd.read_csv(TAXONOMY_FILE)

required_taxonomy_columns = [
    "skill_name_en",
    "description_en",
    "subsector_en",
    "related_job_families_en",
    "needs_review"
]

missing_columns = [
    col for col in required_taxonomy_columns
    if col not in taxonomy.columns
]

if missing_columns:
    raise ValueError(
        f"Taxonomy is missing required columns: {missing_columns}"
    )

taxonomy = taxonomy.dropna(
    subset=["skill_name_en"]
).copy()

taxonomy["skill_name_en"] = (
    taxonomy["skill_name_en"]
    .astype(str)
    .str.strip()
)

taxonomy["description_en"] = (
    taxonomy["description_en"]
    .fillna("")
    .astype(str)
    .str.strip()
)

taxonomy_skills = (
    taxonomy["skill_name_en"]
    .drop_duplicates()
    .tolist()
)

VALID_COMPETENCIES = set(taxonomy_skills)

print()
print("Taxonomy shape:", taxonomy.shape)
print("Number of valid taxonomy skills:", len(taxonomy_skills))


# ============================================================
# 3. TAXONOMY CANDIDATES
# ============================================================
#
# V6 uses taxonomy candidates during final aggregation.
# The frozen taxonomy is the authoritative candidate source.
#
# No competency may be returned unless it exactly matches
# one of these skill names.
# ============================================================

matches = taxonomy.copy()

print("Available taxonomy candidates:", len(matches))


# ============================================================
# 4. SYSTEM PROMPT
# ============================================================

SYSTEM_PROMPT = """
You are an AI system for learning-content tagging and competency mapping.

Analyze educational learning content and return structured JSON.

Your task is to identify:

1. Specific topic tags
2. Saudi Skills Taxonomy competencies
3. Difficulty level
4. Learning objectives
5. Confidence
6. Brief evidence-based notes

IMPORTANT:

- Use only evidence present in the supplied learning content.
- Do not invent information.
- Do not use the filename as evidence.
- Tags must be specific concepts, algorithms, techniques, tools, or procedures.
- Return 3 to 8 tags.
- Competencies must be copied EXACTLY from the provided Saudi Skills Taxonomy.
- Return 1 to 3 competencies.
- Never invent a competency.
- Difficulty must be Beginner, Intermediate, or Advanced.
- Learning objectives must describe what a learner should be able to do after studying the content.
- Return 2 to 4 concise learning objectives.
- Return ONLY valid JSON.
"""


# ============================================================
# 5. CONTENT PROMPT
# ============================================================

def build_prompt(content):

    candidate_skills = matches[
        ["skill_name_en", "description_en"]
    ].to_dict("records")

    skills_text = "\n".join(
        f"- {item['skill_name_en']}: {item['description_en']}"
        for item in candidate_skills
    )

    return f"""
You are analyzing educational learning content.

AVAILABLE SAUDI SKILLS TAXONOMY:

{skills_text}


========================
TASK
========================

Analyze the content below.


========================
TAGS
========================

Return 3 to 8 specific tags.

Use actual concepts, algorithms, techniques, tools, and procedures
found in the content.

Do not invent tags.

Do not use the filename as evidence.

Remove duplicates.


========================
COMPETENCIES
========================

Select 1 to 3 competencies that are genuinely supported by the content.

A competency is valid when the content meaningfully teaches,
explains, demonstrates, or practices the skill.

Do NOT select a competency only because it is generally related
to the topic.

Every competency MUST be copied EXACTLY from the taxonomy list.

Do not invent competency names.


========================
DIFFICULTY
========================

Use only:

Beginner
Intermediate
Advanced

Choose based on the actual technical complexity of the content.


========================
LEARNING OBJECTIVES
========================

Return 2 to 4 concise learning objectives.

Each objective should describe something the learner should be able
to understand, explain, implement, analyze, or apply after studying
the content.

Objectives must be grounded in the actual content.

Do not invent topics that are not present.


========================
CONFIDENCE
========================

Return a number from 0 to 1.


========================
NOTES
========================

Briefly explain the competency selection and difficulty.

Mention concrete evidence from the content.


========================
CONTENT
========================

{content}


========================
OUTPUT
========================

Return ONLY valid JSON.

No Markdown.
No ```json.

{{
  "predicted_tags": [
    "tag 1",
    "tag 2",
    "tag 3"
  ],
  "proposed_competencies": [
    "EXACT taxonomy skill name"
  ],
  "difficulty_level": "Beginner",
  "learning_objectives": [
    "Learning objective 1",
    "Learning objective 2"
  ],
  "confidence": 0.0,
  "notes": "Evidence-based explanation."
}}
"""


# ============================================================
# 6. JSON CLEANING
# ============================================================

def clean_qwen_json(response):

    response = response.strip()

    if response.startswith("```"):
        response = response.replace("```json", "")
        response = response.replace("```", "")
        response = response.strip()

    return json.loads(response)


# ============================================================
# 7. RESULT VALIDATION
# ============================================================

def validate_result(result):

    # -------------------------
    # Tags
    # -------------------------

    tags = result.get(
        "predicted_tags",
        []
    )

    if not isinstance(tags, list):
        tags = []

    tags = [
        str(tag).strip()
        for tag in tags
        if str(tag).strip()
    ]

    tags = tags[:8]


    # -------------------------
    # Competencies
    # -------------------------

    competencies = result.get(
        "proposed_competencies",
        []
    )

    if not isinstance(competencies, list):
        competencies = []

    valid_competencies = []

    for competency in competencies:

        competency = str(
            competency
        ).strip()

        if competency in VALID_COMPETENCIES:

            if competency not in valid_competencies:
                valid_competencies.append(
                    competency
                )

    valid_competencies = valid_competencies[:3]


    # -------------------------
    # Difficulty
    # -------------------------

    difficulty = result.get(
        "difficulty_level",
        "Beginner"
    )

    if difficulty not in [
        "Beginner",
        "Intermediate",
        "Advanced"
    ]:
        difficulty = "Beginner"


    # -------------------------
    # Learning Objectives
    # -------------------------

    learning_objectives = result.get(
        "learning_objectives",
        []
    )

    if not isinstance(
        learning_objectives,
        list
    ):
        learning_objectives = []

    learning_objectives = [
        str(objective).strip()
        for objective in learning_objectives
        if str(objective).strip()
    ]

    # Remove duplicate objectives
    unique_objectives = []
    seen_objectives = set()

    for objective in learning_objectives:

        key = objective.lower()

        if key not in seen_objectives:

            unique_objectives.append(
                objective
            )

            seen_objectives.add(key)

    learning_objectives = unique_objectives[:4]


    # -------------------------
    # Confidence
    # -------------------------

    try:

        confidence = float(
            result.get(
                "confidence",
                0
            )
        )

    except Exception:

        confidence = 0.0

    confidence = max(
        0.0,
        min(1.0, confidence)
    )


    # -------------------------
    # Notes
    # -------------------------

    notes = str(
        result.get(
            "notes",
            ""
        )
    ).strip()


    return {
        "predicted_tags": tags,
        "proposed_competencies": valid_competencies,
        "difficulty_level": difficulty,
        "learning_objectives": learning_objectives,
        "confidence": confidence,
        "notes": notes
    }


# ============================================================
# 8. QWEN SINGLE-CHUNK INFERENCE
# ============================================================

def run_qwen(content):

    messages = [
        {
            "role": "system",
            "content": SYSTEM_PROMPT
        },
        {
            "role": "user",
            "content": build_prompt(content)
        }
    ]

    prompt = tokenizer.apply_chat_template(
        messages,
        tokenize=False,
        add_generation_prompt=True
    )

    inputs = tokenizer(
        prompt,
        return_tensors="pt"
    ).to("cuda")

    with torch.no_grad():

        outputs = model.generate(
            **inputs,
            max_new_tokens=700,
            do_sample=False
        )

    generated_tokens = outputs[
        0
    ][
        inputs["input_ids"].shape[1]:
    ]

    response = tokenizer.decode(
        generated_tokens,
        skip_special_tokens=True
    )

    del inputs
    del outputs
    del generated_tokens

    gc.collect()
    torch.cuda.empty_cache()

    try:

        result = clean_qwen_json(
            response
        )

        return validate_result(
            result
        )

    except Exception as e:

        return {
            "predicted_tags": [],
            "proposed_competencies": [],
            "difficulty_level": "Beginner",
            "learning_objectives": [],
            "confidence": 0.0,
            "notes": (
                "JSON parsing/validation error: "
                f"{str(e)}"
            )
        }


# ============================================================
# 9. FINAL AGGREGATION PROMPT
# ============================================================

def build_aggregation_prompt(
    chunk_results
):

    chunk_evidence = []

    for i, result in enumerate(
        chunk_results,
        start=1
    ):

        chunk_evidence.append(
            f"""
CHUNK {i}

Tags:
{", ".join(result.get("predicted_tags", []))}

Suggested competencies:
{", ".join(result.get("proposed_competencies", []))}

Learning objectives:
{chr(10).join("- " + x for x in result.get("learning_objectives", []))}

Difficulty:
{result.get("difficulty_level", "Beginner")}

Confidence:
{result.get("confidence", 0)}

Evidence:
{result.get("notes", "")}
"""
        )

    evidence_text = "\n".join(
        chunk_evidence
    )

    candidate_skills = matches[
        ["skill_name_en", "description_en"]
    ].to_dict("records")

    skills_text = "\n".join(
        f"- {item['skill_name_en']}: {item['description_en']}"
        for item in candidate_skills
    )

    return f"""
You are the final decision stage of a learning-content tagging system.

Your job is to produce the FINAL result for the complete learning
content using evidence collected from multiple chunks.

The chunk results are evidence. They are not ground truth.

AVAILABLE TAXONOMY CANDIDATES:

{skills_text}


========================
CORE TASK
========================

Select 1 to 3 competencies genuinely supported by the learning
content.

A competency is valid when the content meaningfully teaches,
explains, demonstrates, or practices the skill.

Do NOT select a competency only because it is generally related
to the topic.

Do NOT invent competencies.

Every competency MUST be copied EXACTLY from the taxonomy candidate list.


========================
HOW TO COMBINE CHUNKS
========================

Consider evidence across ALL chunks.

If the same competency is supported by several chunks, this is strong
evidence.

If a competency appears in only one chunk but that chunk contains
clear and specific evidence for the competency, it can still be valid.

Do NOT require a competency to appear in multiple chunks.

Do NOT remove a valid competency simply because another competency
is more central.

A document may legitimately have 2 or 3 competencies.

However, do not fill the 3 slots with weak or unrelated competencies.


========================
IMPORTANT DISTINCTIONS
========================

"Machine Learning (ML)"
is appropriate when the content teaches machine learning algorithms,
models, workflows, training, prediction, classification, regression,
decision trees, KNN, ensemble learning, or related ML techniques.

"Data analytics"
is appropriate when the content teaches data analysis,
interpretation, analytical methods, statistics, metrics, or extracting
insights from data.

"Data preparation"
is appropriate when the content teaches preparing, cleaning,
transforming, structuring, or selecting data for analysis or ML.

"Data processing"
is appropriate when the content teaches manipulating, transforming,
querying, or processing data.

"Database management and configuration"
is appropriate when the content teaches SQL/database interaction,
database operations, querying, joins, or management of databases.

"Natural language Processing (NLP)"
is appropriate when the content teaches language processing,
text processing, sentiment analysis, language understanding,
prompt/language techniques, retrieval over language content,
or related NLP concepts.

"Artificial intelligence, machine learning and deep learning application"
is appropriate when the content meaningfully teaches or applies
AI, machine learning, or deep learning concepts.

"Artificial intelligence application in product development"
is highly specific.

Select it ONLY if the actual evidence explicitly concerns AI
applications in product development or product-related contexts.

AI, ML, NLP, RAG, embeddings, or prompting by themselves are NOT
evidence for product development.


========================
DECISION RULE
========================

For every candidate competency, ask:

1. What evidence in the chunks supports it?
2. Is that evidence specific and meaningful?
3. Is the competency actually taught or practiced?

If YES, keep it.

If the connection is only broad or indirect, remove it.

After evaluating all candidates, return the strongest 1 to 3
competencies supported by the evidence.


========================
TAGS
========================

Return 3 to 8 specific tags.

Use actual concepts, algorithms, techniques, tools, and procedures
found in the chunks.

Remove duplicates.

Do not invent tags.

Do not use the filename as evidence.


========================
DIFFICULTY
========================

Use only:

Beginner
Intermediate
Advanced

Choose based on the actual technical complexity of the content.


========================
LEARNING OBJECTIVES
========================

Return 2 to 4 concise learning objectives.

Objectives must be grounded in the actual content.

Describe what a learner should be able to understand, explain,
implement, analyze, or apply after studying the content.

Do not introduce topics that are not taught in the content.


========================
CONFIDENCE
========================

Return a number from 0 to 1 representing confidence in the final
classification.


========================
NOTES
========================

Briefly explain why the final competencies were selected.

Mention concrete evidence from the chunks.


========================
CHUNK EVIDENCE
========================

{evidence_text}


========================
OUTPUT
========================

Return ONLY valid JSON.

No Markdown.
No ```json.

{{
  "predicted_tags": [
    "tag 1",
    "tag 2",
    "tag 3"
  ],
  "proposed_competencies": [
    "EXACT taxonomy skill name"
  ],
  "difficulty_level": "Beginner",
  "learning_objectives": [
    "Learning objective 1",
    "Learning objective 2"
  ],
  "confidence": 0.0,
  "notes": "Evidence-based explanation."
}}
"""


# ============================================================
# 10. CONTENT EXTRACTION
# ============================================================

def extract_pptx(path):

    prs = Presentation(path)

    texts = []

    for slide_num, slide in enumerate(
        prs.slides,
        start=1
    ):

        slide_text = []

        for shape in slide.shapes:

            if hasattr(shape, "text"):

                text = shape.text.strip()

                if text:
                    slide_text.append(text)

        if slide_text:

            texts.append(
                f"[Slide {slide_num}]\n"
                + "\n".join(slide_text)
            )

    return "\n\n".join(texts)


def extract_ipynb(path):

    with open(
        path,
        "r",
        encoding="utf-8"
    ) as f:

        notebook = json.load(f)

    texts = []

    for cell_num, cell in enumerate(
        notebook.get("cells", []),
        start=1
    ):

        cell_type = cell.get(
            "cell_type",
            ""
        )

        source = "".join(
            cell.get("source", [])
        )

        if not source.strip():
            continue

        if cell_type == "markdown":

            texts.append(
                f"[Markdown Cell {cell_num}]\n"
                f"{source}"
            )

        elif cell_type == "code":

            texts.append(
                f"[Code Cell {cell_num}]\n"
                f"{source}"
            )

    return "\n\n".join(texts)


def extract_xlsx(path):

    workbook = load_workbook(
        path,
        read_only=True,
        data_only=True
    )

    texts = []

    for sheet in workbook.worksheets:

        sheet_rows = []

        for row in sheet.iter_rows(
            values_only=True
        ):

            values = [
                str(value).strip()
                for value in row
                if value is not None
            ]

            if values:

                sheet_rows.append(
                    " | ".join(values)
                )

        if sheet_rows:

            texts.append(
                f"[Sheet: {sheet.title}]\n"
                + "\n".join(sheet_rows)
            )

    return "\n\n".join(texts)


def extract_text_file(path):

    with open(
        path,
        "r",
        encoding="utf-8",
        errors="ignore"
    ) as f:

        return f.read()


def extract_content(path):

    extension = Path(
        path
    ).suffix.lower()

    if extension == ".pptx":
        return extract_pptx(path)

    elif extension == ".ipynb":
        return extract_ipynb(path)

    elif extension == ".xlsx":
        return extract_xlsx(path)

    elif extension in [
        ".md",
        ".txt"
    ]:
        return extract_text_file(path)

    else:
        return None


# ============================================================
# 11. FILE DISCOVERY
# ============================================================

SUPPORTED_EXTENSIONS = [
    ".pptx",
    ".ipynb",
    ".xlsx",
    ".md",
    ".txt"
]

files = sorted(
    [
        str(Path("evaluation_data") / f)
        for f in os.listdir("evaluation_data")
        if Path(f).suffix.lower()
        in SUPPORTED_EXTENSIONS
    ]
)

print()
print("=" * 80)
print("FILES FOUND")
print("=" * 80)
print("Files found:", len(files))

for file in files:

    print(
        f"{file}"
    )


# ============================================================
# 12. CHUNKING
# ============================================================

def split_content(
    content,
    chunk_size=6000,
    overlap=500
):

    chunks = []

    start = 0

    content_length = len(
        content
    )

    while start < content_length:

        end = min(
            start + chunk_size,
            content_length
        )

        chunk = content[
            start:end
        ].strip()

        if chunk:
            chunks.append(
                chunk
            )

        if end >= content_length:
            break

        start = end - overlap

    return chunks


# ============================================================
# 13. CHUNKED INFERENCE
# ============================================================

def run_qwen_chunked(
    content,
    chunk_size=6000,
    overlap=500
):

    chunks = split_content(
        content,
        chunk_size=chunk_size,
        overlap=overlap
    )

    chunk_results = []

    print(
        f"Total chunks: {len(chunks)}"
    )

    for i, chunk in enumerate(
        chunks,
        start=1
    ):

        print(
            f"\n{'=' * 60}"
        )

        print(
            f"Processing chunk "
            f"{i}/{len(chunks)}"
        )

        print(
            f"Characters: {len(chunk)}"
        )

        try:

            result = run_qwen(
                chunk
            )

            chunk_results.append(
                result
            )

            print(
                "Tags:",
                result["predicted_tags"]
            )

            print(
                "Competencies:",
                result["proposed_competencies"]
            )

            print(
                "Difficulty:",
                result["difficulty_level"]
            )

            print(
                "Learning objectives:",
                result["learning_objectives"]
            )

            print(
                "Confidence:",
                result["confidence"]
            )

        except Exception as e:

            print(
                f"Error in chunk {i}: {e}"
            )

            chunk_results.append(
                {
                    "predicted_tags": [],
                    "proposed_competencies": [],
                    "difficulty_level": "Beginner",
                    "learning_objectives": [],
                    "confidence": 0.0,
                    "notes": (
                        f"Chunk error: {str(e)}"
                    )
                }
            )

        gc.collect()

        torch.cuda.empty_cache()

    return chunk_results


# ============================================================
# 14. SIMPLE CHUNK AGGREGATION
# ============================================================

def aggregate_chunk_results(
    chunk_results
):

    # -------------------------
    # Tags
    # -------------------------

    all_tags = []

    for result in chunk_results:

        all_tags.extend(
            result.get(
                "predicted_tags",
                []
            )
        )

    unique_tags = []

    seen = set()

    for tag in all_tags:

        tag_clean = tag.strip()

        key = tag_clean.lower()

        if (
            tag_clean
            and key not in seen
        ):

            unique_tags.append(
                tag_clean
            )

            seen.add(key)


    # -------------------------
    # Competencies
    # -------------------------

    all_competencies = []

    for result in chunk_results:

        all_competencies.extend(
            result.get(
                "proposed_competencies",
                []
            )
        )

    unique_competencies = []

    seen_comp = set()

    for competency in all_competencies:

        competency_clean = (
            competency.strip()
        )

        if (
            competency_clean
            and competency_clean
            in VALID_COMPETENCIES
            and competency_clean
            not in seen_comp
        ):

            unique_competencies.append(
                competency_clean
            )

            seen_comp.add(
                competency_clean
            )


    # -------------------------
    # Learning Objectives
    # -------------------------

    all_objectives = []

    for result in chunk_results:

        all_objectives.extend(
            result.get(
                "learning_objectives",
                []
            )
        )

    unique_objectives = []

    seen_objectives = set()

    for objective in all_objectives:

        objective_clean = (
            objective.strip()
        )

        key = objective_clean.lower()

        if (
            objective_clean
            and key not in seen_objectives
        ):

            unique_objectives.append(
                objective_clean
            )

            seen_objectives.add(
                key
            )


    # -------------------------
    # Difficulty votes
    # -------------------------

    difficulty_counts = {
        "Beginner": 0,
        "Intermediate": 0,
        "Advanced": 0
    }

    for result in chunk_results:

        difficulty = result.get(
            "difficulty_level"
        )

        if difficulty in difficulty_counts:

            difficulty_counts[
                difficulty
            ] += 1

    final_difficulty = max(
        difficulty_counts,
        key=difficulty_counts.get
    )


    # -------------------------
    # Confidence
    # -------------------------

    confidences = [
        float(
            r.get(
                "confidence",
                0
            )
        )
        for r in chunk_results
        if r.get("confidence")
        is not None
    ]

    final_confidence = (
        sum(confidences)
        / len(confidences)
        if confidences
        else 0.0
    )


    # -------------------------
    # Notes
    # -------------------------

    notes = []

    for i, result in enumerate(
        chunk_results,
        start=1
    ):

        note = result.get(
            "notes",
            ""
        ).strip()

        if note:

            notes.append(
                f"Chunk {i}: {note}"
            )


    return {
        "predicted_tags":
            unique_tags[:8],

        "proposed_competencies":
            unique_competencies[:3],

        "difficulty_level":
            final_difficulty,

        "learning_objectives":
            unique_objectives[:4],

        "confidence":
            round(
                final_confidence,
                2
            ),

        "notes":
            " ".join(notes)
    }


# ============================================================
# 15. FINAL QWEN AGGREGATOR
# ============================================================

def run_final_aggregator(
    chunk_results
):

    messages = [
        {
            "role": "system",
            "content": SYSTEM_PROMPT
        },
        {
            "role": "user",
            "content":
                build_aggregation_prompt(
                    chunk_results
                )
        }
    ]

    prompt = tokenizer.apply_chat_template(
        messages,
        tokenize=False,
        add_generation_prompt=True
    )

    inputs = tokenizer(
        prompt,
        return_tensors="pt"
    ).to("cuda")

    with torch.no_grad():

        outputs = model.generate(
            **inputs,
            max_new_tokens=500,
            do_sample=False
        )

    generated_tokens = outputs[
        0
    ][
        inputs["input_ids"].shape[1]:
    ]

    response = tokenizer.decode(
        generated_tokens,
        skip_special_tokens=True
    )

    del inputs
    del outputs
    del generated_tokens

    gc.collect()

    torch.cuda.empty_cache()

    try:

        result = clean_qwen_json(
            response
        )

        return validate_result(
            result
        )

    except Exception as e:

        return {
            "predicted_tags": [],
            "proposed_competencies": [],
            "difficulty_level": "Beginner",
            "learning_objectives": [],
            "confidence": 0.0,
            "notes": (
                "Final aggregation error: "
                f"{str(e)}"
            )
        }


# ============================================================
# 16. PROCESS ALL EVALUATION FILES
# ============================================================

chunk_results_by_file = {}

print()
print("=" * 80)
print("STARTING V6 CHUNK PROCESSING")
print("=" * 80)

for file_name in files:

    print()
    print("=" * 80)
    print("FILE:", file_name)
    print("=" * 80)

    try:

        content = extract_content(
            file_name
        )

        if not content:

            print(
                "WARNING: No content extracted."
            )

            chunk_results_by_file[
                file_name
            ] = []

            continue

        print(
            "Characters:",
            len(content)
        )

        chunk_results = (
            run_qwen_chunked(
                content
            )
        )

        chunk_results_by_file[
            file_name
        ] = chunk_results

    except Exception as e:

        print(
            f"File processing error: {e}"
        )

        chunk_results_by_file[
            file_name
        ] = []


# ============================================================
# 17. FINAL AGGREGATION FOR ALL FILES
# ============================================================

final_results_v6 = []

print()
print("=" * 80)
print("FINAL AGGREGATION V6")
print("=" * 80)

for file_name, chunk_results in (
    chunk_results_by_file.items()
):

    print(
        f"\n{'=' * 70}"
    )

    print(
        f"FINAL AGGREGATION V6: "
        f"{file_name}"
    )

    try:

        if not chunk_results:

            raise ValueError(
                "No chunk results available."
            )

        result = run_final_aggregator(
            chunk_results
        )

        result["file_name"] = (
            file_name
        )

        final_results_v6.append(
            result
        )

        print(
            "Tags:",
            result[
                "predicted_tags"
            ]
        )

        print(
            "Competencies:",
            result[
                "proposed_competencies"
            ]
        )

        print(
            "Difficulty:",
            result[
                "difficulty_level"
            ]
        )

        print(
            "Learning objectives:",
            result[
                "learning_objectives"
            ]
        )

        print(
            "Confidence:",
            result[
                "confidence"
            ]
        )

        print(
            "Notes:",
            result[
                "notes"
            ]
        )

    except Exception as e:

        print(
            f"Error: {e}"
        )

        final_results_v6.append(
            {
                "file_name":
                    file_name,

                "predicted_tags":
                    [],

                "proposed_competencies":
                    [],

                "difficulty_level":
                    "Beginner",

                "learning_objectives":
                    [],

                "confidence":
                    0.0,

                "notes":
                    f"Aggregation error: {str(e)}"
            }
        )


# ============================================================
# 18. FINAL DATAFRAME
# ============================================================

results_df_v6 = pd.DataFrame(
    final_results_v6
)

output_columns = [
    "file_name",
    "predicted_tags",
    "proposed_competencies",
    "difficulty_level",
    "learning_objectives",
    "confidence",
    "notes"
]

results_df_v6 = results_df_v6[
    output_columns
]


# ============================================================
# 19. SAVE V6 OUTPUT
# ============================================================

OUTPUT_FILE = (
    "qwen7b_final_predictions_v6.csv"
)

results_df_v6.to_csv(
    OUTPUT_FILE,
    index=False,
    encoding="utf-8-sig"
)

print()
print("=" * 80)
print("V6 COMPLETE")
print("=" * 80)

print(
    "Output:",
    OUTPUT_FILE
)

print(
    "Rows:",
    len(results_df_v6)
)

print()
print(
    results_df_v6.to_string(
        index=False
    )
)

print()
print("=" * 80)
print("REQUIRED OUTPUT FIELDS")
print("=" * 80)

print(
    results_df_v6.columns.tolist()
)

print()
print(
    "V6 model pipeline finished."
)