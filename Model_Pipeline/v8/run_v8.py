"""
V8 inference runner.

Reads official evaluation files, extracts text, runs the
V8 Qwen + Hybrid E5 pipeline, and saves predictions.

Ground-truth labels are NOT used during inference.
"""

import csv
import json
import sys
from pathlib import Path

import nbformat
from openpyxl import load_workbook
from pptx import Presentation

from model_v8 import process_document


# ============================================================
# PATHS
# ============================================================

ROOT_DIR = Path(__file__).resolve().parents[2]

EVALUATION_DIR = (
    ROOT_DIR
    / "Model"
    / "evaluation_data"
)

OUTPUT_DIR = (
    ROOT_DIR
    / "Model"
    / "v8"
    / "results"
)

OUTPUT_JSON = OUTPUT_DIR / "v8_predictions.json"
OUTPUT_CSV = OUTPUT_DIR / "v8_predictions.csv"


# ============================================================
# CONTENT EXTRACTION
# ============================================================

def extract_pptx(path):
    presentation = Presentation(path)

    parts = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1,
    ):
        slide_parts = []

        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text = shape.text.strip()

                if text:
                    slide_parts.append(text)

        if slide_parts:
            parts.append(
                f"[SLIDE {slide_number}]\n"
                + "\n".join(slide_parts)
            )

    return "\n\n".join(parts)


def extract_ipynb(path):
    notebook = nbformat.read(
        path,
        as_version=4,
    )

    parts = []

    for cell_number, cell in enumerate(
        notebook.cells,
        start=1,
    ):
        cell_type = cell.cell_type
        source = cell.source.strip()

        if not source:
            continue

        parts.append(
            f"[{cell_type.upper()} CELL {cell_number}]\n"
            f"{source}"
        )

    return "\n\n".join(parts)


def extract_xlsx(path):
    workbook = load_workbook(
        path,
        read_only=True,
        data_only=True,
    )

    parts = []

    for worksheet in workbook.worksheets:

        rows = []

        for row in worksheet.iter_rows(
            values_only=True
        ):
            values = []

            for value in row:
                if value is not None:
                    values.append(str(value))

            if values:
                rows.append(" | ".join(values))

        if rows:
            parts.append(
                f"[SHEET: {worksheet.title}]\n"
                + "\n".join(rows)
            )

    return "\n\n".join(parts)


def extract_text_file(path):
    return path.read_text(
        encoding="utf-8",
        errors="ignore",
    )


def extract_content(path):
    suffix = path.suffix.lower()

    if suffix == ".pptx":
        return extract_pptx(path)

    if suffix == ".ipynb":
        return extract_ipynb(path)

    if suffix == ".xlsx":
        return extract_xlsx(path)

    if suffix in {
        ".md",
        ".txt",
        ".csv",
    }:
        return extract_text_file(path)

    raise ValueError(
        f"Unsupported file type: {path.suffix}"
    )


# ============================================================
# OUTPUT HELPERS
# ============================================================

def flatten_prediction(file_name, result):

    return {
        "file_name": file_name,
        "predicted_tags": json.dumps(
            result["predicted_tags"],
            ensure_ascii=False,
        ),
        "proposed_competencies": json.dumps(
            result["proposed_competencies"],
            ensure_ascii=False,
        ),
        "difficulty_level": result[
            "difficulty_level"
        ],
        "confidence": result[
            "confidence"
        ],
        "notes": result["notes"],
        "chunk_count": result[
            "chunk_count"
        ],
        "retrieval_candidates": json.dumps(
            result["retrieval_candidates"],
            ensure_ascii=False,
        ),
    }


# ============================================================
# RUN ONE FILE
# ============================================================

def run_one(path):

    print("\n" + "=" * 80)
    print(f"V8 FILE: {path.name}")
    print("=" * 80)

    content = extract_content(path)

    if not content.strip():
        raise ValueError(
            f"No text extracted from {path.name}"
        )

    print(
        f"Extracted characters: {len(content):,}"
    )

    result = process_document(content)

    print("\nFinal tags:")
    for tag in result["predicted_tags"]:
        print(f"  - {tag}")

    print("\nFinal competencies:")
    for competency in result[
        "proposed_competencies"
    ]:
        print(f"  - {competency}")

    print(
        f"\nDifficulty: "
        f"{result['difficulty_level']}"
    )

    print(
        f"Confidence: "
        f"{result['confidence']:.3f}"
    )

    return result


# ============================================================
# MAIN
# ============================================================

def main():

    OUTPUT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    if len(sys.argv) > 1:
        requested = [
            Path(x)
            for x in sys.argv[1:]
        ]

        files = []

        for path in requested:

            if path.exists():
                files.append(path)

            else:
                candidate = (
                    EVALUATION_DIR / path.name
                )

                if candidate.exists():
                    files.append(candidate)

                else:
                    raise FileNotFoundError(
                        f"File not found: {path}"
                    )

    else:
        files = sorted(
            EVALUATION_DIR.iterdir()
        )

        files = [
            path
            for path in files
            if path.suffix.lower()
            in {
                ".pptx",
                ".ipynb",
                ".xlsx",
                ".md",
                ".txt",
                ".csv",
            }
        ]

    print(
        f"Files selected: {len(files)}"
    )

    predictions = {}

    rows = []

    for index, path in enumerate(
        files,
        start=1,
    ):

        print(
            f"\nPROCESSING "
            f"{index}/{len(files)}: "
            f"{path.name}"
        )

        try:

            result = run_one(path)

            predictions[path.name] = result

            rows.append(
                flatten_prediction(
                    path.name,
                    result,
                )
            )

        except Exception as exc:

            print(
                f"ERROR processing "
                f"{path.name}: {exc}"
            )

            predictions[path.name] = {
                "error": str(exc)
            }

            rows.append({
                "file_name": path.name,
                "predicted_tags": "[]",
                "proposed_competencies": "[]",
                "difficulty_level": "",
                "confidence": 0.0,
                "notes": f"ERROR: {exc}",
                "chunk_count": 0,
                "retrieval_candidates": "[]",
            })

    OUTPUT_JSON.write_text(
        json.dumps(
            predictions,
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )

    fieldnames = [
        "file_name",
        "predicted_tags",
        "proposed_competencies",
        "difficulty_level",
        "confidence",
        "notes",
        "chunk_count",
        "retrieval_candidates",
    ]

    with OUTPUT_CSV.open(
        "w",
        encoding="utf-8",
        newline="",
    ) as f:

        writer = csv.DictWriter(
            f,
            fieldnames=fieldnames,
        )

        writer.writeheader()
        writer.writerows(rows)

    print("\n" + "=" * 80)
    print("V8 RUN COMPLETE")
    print("=" * 80)

    print(f"JSON: {OUTPUT_JSON}")
    print(f"CSV : {OUTPUT_CSV}")


if __name__ == "__main__":
    main()
