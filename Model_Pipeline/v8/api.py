import json
import sys
import tempfile
from pathlib import Path

from fastapi import FastAPI, File, UploadFile, HTTPException
from fastapi.concurrency import run_in_threadpool
from pptx import Presentation
from openpyxl import load_workbook
import nbformat

sys.path.insert(0, str(Path(__file__).resolve().parent))

from model_v8 import process_document


app = FastAPI(
    title="BeamData Content Tagging & Competency Mapping",
    version="8.4",
)


def extract_pptx(path):
    presentation = Presentation(path)
    parts = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1,
    ):
        slide_parts = []

        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text = shape.text.strip()
                if text:
                    slide_parts.append(text)

        if slide_parts:
            parts.append(
                f"[SLIDE {slide_number}]\n"
                + "\n".join(slide_parts)
            )

    return "\n\n".join(parts)


def extract_ipynb(path):
    notebook = nbformat.read(
        path,
        as_version=4,
    )

    parts = []

    for cell_number, cell in enumerate(
        notebook.cells,
        start=1,
    ):
        source = cell.source.strip()

        if not source:
            continue

        parts.append(
            f"[{cell.cell_type.upper()} CELL {cell_number}]\n"
            f"{source}"
        )

    return "\n\n".join(parts)


def extract_xlsx(path):
    workbook = load_workbook(
        path,
        read_only=True,
        data_only=True,
    )

    parts = []

    for worksheet in workbook.worksheets:
        rows = []

        for row in worksheet.iter_rows(
            values_only=True
        ):
            values = [
                str(value)
                for value in row
                if value is not None
            ]

            if values:
                rows.append(" | ".join(values))

        if rows:
            parts.append(
                f"[SHEET: {worksheet.title}]\n"
                + "\n".join(rows)
            )

    return "\n\n".join(parts)


def extract_content(path):
    suffix = path.suffix.lower()

    if suffix == ".pptx":
        return extract_pptx(path)

    if suffix == ".ipynb":
        return extract_ipynb(path)

    if suffix == ".xlsx":
        return extract_xlsx(path)

    if suffix in {".md", ".txt", ".csv"}:
        return path.read_text(
            encoding="utf-8",
            errors="ignore",
        )

    raise ValueError(
        f"Unsupported file type: {suffix}"
    )


@app.get("/health")
def health():
    return {
        "status": "ok",
        "service": "beamdata-v8-api",
        "version": "8.4",
    }


@app.post("/predict")
async def predict(
    file: UploadFile = File(...)
):
    suffix = Path(file.filename).suffix.lower()

    supported = {
        ".pptx",
        ".ipynb",
        ".xlsx",
        ".md",
        ".txt",
        ".csv",
    }

    if suffix not in supported:
        raise HTTPException(
            status_code=400,
            detail=f"Unsupported file type: {suffix}",
        )

    data = await file.read()

    if not data:
        raise HTTPException(
            status_code=400,
            detail="Uploaded file is empty.",
        )

    temp_path = None

    try:
        with tempfile.NamedTemporaryFile(
            suffix=suffix,
            delete=False,
        ) as temp:
            temp.write(data)
            temp_path = Path(temp.name)

        content = extract_content(temp_path)

        if not content.strip():
            raise HTTPException(
                status_code=400,
                detail="No text could be extracted.",
            )

        result = await run_in_threadpool(process_document, content)

        return {
            "file_name": file.filename,
            **result,
        }

    except HTTPException:
        raise

    except Exception as exc:
        raise HTTPException(
            status_code=500,
            detail=str(exc),
        )

    finally:
        if temp_path and temp_path.exists():
            temp_path.unlink()


if __name__ == "__main__":
    import uvicorn

    uvicorn.run(
        app,
        host="0.0.0.0",
        port=8000,
    )
